import tempfile, os, importlib.util, io, json
from pathlib import Path
from werkzeug.security import generate_password_hash

def run():
 with tempfile.TemporaryDirectory() as tmp:
  os.environ.update(DOCUMENT_ROOT=tmp+'/documents',STATE_ROOT=tmp+'/state',SESSION_SECRET='test-only-secret',PASSWORD_HASH=generate_password_hash('test-password'))
  spec=importlib.util.spec_from_file_location('desk',Path(__file__).with_name('app.py'));m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  c=m.app.test_client()
  assert c.get('/api/state').status_code==401
  assert c.get('/static/style.css').status_code==200
  assert c.post('/login',data={'password':'wrong'}).status_code==401
  assert c.post('/login',data={'password':'test-password'}).status_code==302
  with c.session_transaction() as s:h={'X-CSRF-Token':s['csrf']}
  assert c.post('/api/upload').status_code==403
  assert c.post('/api/upload',headers=h,data={'file':(io.BytesIO(b'MZ'),'bad.exe')}).status_code==400
  payload=b'Issuer: Test Electric\nAmount: $42'
  r=c.post('/api/upload',headers=h,data={'file':(io.BytesIO(payload),'bill.txt')});assert r.status_code==200,r.data
  d=r.json;assert 'Test Electric' in d['text'];id=d['id']
  assert c.post('/api/upload',headers=h,data={'file':(io.BytesIO(payload),'again.txt')}).status_code==409
  assert c.put('/api/templates',headers=h,json={'name':'Test bill','folder':'bills','fields':['Issuer','Amount'],'required':['Amount']}).status_code==200
  values=dict(type='Test bill',fields={'Issuer':'Test Electric'},filename='electric.txt',folder='bills',text=d['text'])
  assert c.put('/api/docs/'+id,headers=h,json=values).status_code==200
  assert c.post('/api/docs/'+id+'/approve',headers=h,json={}).status_code==400
  values['fields']['Amount']='$42';c.put('/api/docs/'+id,headers=h,json=values)
  r=c.post('/api/docs/'+id+'/approve',headers=h,json={});assert r.status_code==200,r.data
  path=Path(tmp)/'documents'/r.json['path'];assert path.read_bytes()==payload
  assert (path.parent/'record.json').exists() and (path.parent/'extracted.txt').exists()
  assert c.get('/api/docs/'+id+'/original').data==payload
  assert b'$42' in c.get('/api/docs/'+id+'/note').data
  assert c.put('/api/docs/'+id,headers=h,json=values).status_code==200
  assert path.read_bytes()==payload
  assert c.post('/api/docs/'+id+'/approve',headers=h,json={}).status_code==200
  # Real parser smoke checks: native PDF, scanned PDF, image, Office and saved email.
  import fitz
  from PIL import Image, ImageDraw, ImageFont
  im=Image.new('RGB',(1000,200),'white');ImageDraw.Draw(im).text((20,40),'RECEIPT TOTAL 123.45',font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',40),fill='black')
  im.save(Path(tmp)/'receipt.png');assert '123.45' in m.extract(Path(tmp)/'receipt.png')
  pdf=fitz.open();page=pdf.new_page();page.insert_text((40,40),'Insurance Policy 12345');pdf.save(Path(tmp)/'native.pdf');pdf.close();assert '12345' in m.extract(Path(tmp)/'native.pdf')
  pdf=fitz.open();page=pdf.new_page(width=1000,height=200);page.insert_image(page.rect,filename=str(Path(tmp)/'receipt.png'));pdf.save(Path(tmp)/'scan.pdf');pdf.close();assert '123.45' in m.extract(Path(tmp)/'scan.pdf')
  from docx import Document
  doc=Document();doc.add_paragraph('Warranty receipt');doc.save(Path(tmp)/'word.docx');assert 'Warranty' in m.extract(Path(tmp)/'word.docx')
  from openpyxl import Workbook
  wb=Workbook();wb.active.append(['Receipt',123]);wb.save(Path(tmp)/'sheet.xlsx');assert '123' in m.extract(Path(tmp)/'sheet.xlsx')
  mail=Path(tmp)/'mail.eml';mail.write_text('From: example@example.com\nSubject: Test email\nContent-Type: text/plain\n\nInsurance renewal');assert 'Insurance renewal' in m.extract(mail)
  print('PASS: authentication, CSRF, executable rejection, text intake, duplicates, templates, required fields, archive, original download, Safenote draft, archived metadata editing, PDF/image OCR, DOCX, XLSX, EML')
if __name__=='__main__':run()
