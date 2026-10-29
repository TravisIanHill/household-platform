import os, json, sqlite3, uuid, re, secrets, hashlib, zipfile
from pathlib import Path
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from io import BytesIO
from flask import Flask, request, session, jsonify, render_template, send_file, abort
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash

ROOT=Path(os.getenv('DOCUMENT_ROOT','/documents'))
STATE=Path(os.getenv('STATE_ROOT','/state'))
ROOT.mkdir(parents=True,exist_ok=True); STATE.mkdir(parents=True,exist_ok=True)
for name in ['inbox','archive']: (ROOT/name).mkdir(exist_ok=True)
app=Flask(__name__, root_path=str(Path(__file__).resolve().parent))
app.config['DOCUMENT_ROOT']=str(ROOT)
app.secret_key=os.environ['SESSION_SECRET']
PASSWORD_HASH=os.environ['PASSWORD_HASH']
app.config.update(MAX_CONTENT_LENGTH=40*1024*1024,SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Strict',SESSION_COOKIE_SECURE=os.getenv('HTTPS','0')=='1',PERMANENT_SESSION_LIFETIME=1800)
DEFAULTS=[{'name':n,'fields':f,'folder':n.lower().replace(' ','-')} for n,f in [('Bill',['Issuer','Document date','Due date','Amount','Account reference']),('Statement',['Issuer','Period','Account reference']),('Major purchase',['Retailer','Purchase date','Item','Amount','Warranty expiry']),('Insurance',['Provider','Policy reference','Coverage start','Coverage end']),('Email',['From','To','Subject','Date'])]]
READERS={'.pdf':'PDF text / scanned PDF OCR','.png':'Image OCR','.jpg':'Image OCR','.jpeg':'Image OCR','.tif':'Image OCR','.tiff':'Image OCR','.webp':'Image OCR','.txt':'Text','.csv':'Text','.eml':'Email','.docx':'Word text','.xlsx':'Excel cells'}

MEDIA_EXTS={'.mov','.mp4','.m4v','.mkv','.avi','.webm','.mp3','.m4a','.wav','.flac','.m4b'}
READERS.update({ext:'Media original retention / manual review' for ext in MEDIA_EXTS})
app.config['READERS']=READERS
app.config['MEDIA_EXTS']=MEDIA_EXTS

def db():
    c=sqlite3.connect(STATE/'desk.sqlite'); c.row_factory=sqlite3.Row; return c
with db() as c:
    c.execute('CREATE TABLE IF NOT EXISTS docs (id TEXT PRIMARY KEY, data TEXT NOT NULL)')
    c.execute('CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, data TEXT NOT NULL)')
    c.execute('INSERT OR IGNORE INTO settings VALUES (?,?)',('templates',json.dumps(DEFAULTS)))

def templates():
    with db() as c:return json.loads(c.execute("SELECT data FROM settings WHERE key='templates'").fetchone()[0])
def getdoc(id):
    with db() as c:r=c.execute('SELECT data FROM docs WHERE id=?',(id,)).fetchone()
    if not r:abort(404)
    return json.loads(r[0])
def save(d):
    expected=d.pop('_expected_version',None)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT data FROM docs WHERE id=?',(d['id'],)).fetchone()
        current=json.loads(row[0]) if row else None
        if expected is not None and current and current.get('version',0)!=expected:abort(409,description='Record changed. Reload before saving.')
        if current:d['version']=current.get('version',0)+1
        c.execute('INSERT OR REPLACE INTO docs VALUES (?,?)',(d['id'],json.dumps(d)))
    if d.get('status')=='archived':
        parent=confined(d['path']).parent
        for name,value in [('record.json',json.dumps(d,indent=2)),('extracted.txt',d.get('text',''))]:
            target=parent/name;temp=parent/(name+'.'+uuid.uuid4().hex+'.tmp');temp.write_text(value);os.replace(temp,target)

def stamp():return datetime.now(timezone.utc).isoformat()
def confined(relative):
    p=(ROOT/relative).resolve()
    if ROOT.resolve() not in p.parents:abort(400)
    return p

def extract(p):
    ext=p.suffix.lower()
    if ext in MEDIA_EXTS:return 'Media original retained. Review title, date, source and purpose manually. Video/audio transcription and technical metadata extraction are not implemented yet.'
    if ext in ['.txt','.csv']:return p.read_text(errors='replace')[:250000]
    if ext=='.eml':
        m=BytesParser(policy=policy.default).parsebytes(p.read_bytes()); b=m.get_body(preferencelist=('plain',))
        return '\n'.join(f'{k}: {m.get(k,"")}' for k in ['From','To','Subject','Date'])+'\n\n'+(b.get_content() if b else 'No plain-text email body; review original email.')[:250000]
    if ext in ['.docx','.xlsx']:
        with zipfile.ZipFile(p) as z:
            if sum(i.file_size for i in z.infolist())>80*1024*1024:raise ValueError('Expanded Office file is too large')
        if ext=='.docx':
            from docx import Document
            doc=Document(p);return '\n'.join([x.text for x in doc.paragraphs]+[' | '.join(cell.text for cell in row.cells) for t in doc.tables for row in t.rows])[:250000]
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True,data_only=True); out=[]
        try:
            for sheet in wb:
                out.append(sheet.title)
                for i,row in enumerate(sheet.iter_rows(values_only=True)):
                    if i>=2000:break
                    out.append(' | '.join(str(x) if x is not None else '' for x in row[:40]))
            return '\n'.join(out)[:250000]
        finally:wb.close()
    import pytesseract
    from PIL import Image
    if ext!='.pdf':
        with Image.open(p) as im: return pytesseract.image_to_string(im,timeout=45)[:250000]
    from ocr_reader import extract_pdf
    return extract_pdf(p)

@app.before_request
def guard():
    if request.path=='/api/upload':request.max_content_length=513*1024*1024
    if request.path=='/api/catalog/import':request.max_content_length=256*1024*1024
    if request.path=='/login' or request.path.startswith('/static/'):return
    if not session.get('signed_in'):
        if request.path.startswith('/api/'):abort(401)
        return render_template('login.html')
    user=app.extensions['household']['user']()
    if not user:
        session.clear();abort(401)
    if request.method in ['POST','PUT','DELETE']:
        if not secrets.compare_digest(request.headers.get('X-CSRF-Token',''),session.get('csrf','')):abort(403)

@app.after_request
def headers(r):
    r.headers['X-Content-Type-Options']='nosniff';r.headers['X-Frame-Options']='SAMEORIGIN';r.headers['Cache-Control']='no-store'
    r.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' blob:; frame-src 'self'; object-src 'none'; frame-ancestors 'self'; base-uri 'none'; form-action 'self'"
    return r

@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='GET':return render_template('login.html')
    username=request.form.get('username','').strip()
    password=request.form.get('password','')
    with db() as c:
        count=c.execute('SELECT count(*) FROM users').fetchone()[0]
        account=c.execute('SELECT * FROM users WHERE username=? AND active=1',(username,)).fetchone() if username else None
    legacy=not username and (count==0 or not any_named_admin()) and check_password_hash(PASSWORD_HASH,password)
    if not legacy and not (account and check_password_hash(account['password'],password)):
        return render_template('login.html',error='Sign-in details not recognized'),401
    session.clear();session.update(signed_in=True,csrf=secrets.token_hex(32));session.permanent=True
    if account:session['user_id']=account['id']
    from flask import redirect
    return redirect('/')
def any_named_admin():
    with db() as c:return bool(c.execute("SELECT 1 FROM users WHERE role='admin' AND active=1").fetchone())
@app.get('/')
def home():return render_template('finder.html',csrf=session['csrf'])
@app.get('/intake')
def intake_page():return render_template('index.html',csrf=session['csrf'])
@app.post('/api/logout')
def logout():session.clear();return jsonify(ok=True)
@app.get('/api/state')
def state():
    with db() as c:docs=[json.loads(r[0]) for r in c.execute('SELECT data FROM docs')]
    docs=[d for d in docs if app.extensions['household']['visible'](d)]
    for d in docs:
        if 'size_bytes' not in d:
            try:d['size_bytes']=confined(d['path']).stat().st_size
            except (OSError, ValueError):d['size_bytes']=None
        d.setdefault('source_modified',None)
    for d in docs:
        d['can_edit']=app.extensions['household']['permission'](d)=='edit'
        d['can_share']=app.extensions['household']['user']()['role']=='admin' or d.get('owner_id')==app.extensions['household']['user']()['id']
    organization=app.extensions['household']['organization']() if app.extensions['household']['user']()['role']=='admin' else dict(folders=[],tags=[],collections=[])
    for d in docs:
        folder=d.get('folder','unclassified')
        organization['folders']+=['/'.join(folder.split('/')[:i]) for i in range(1,len(folder.split('/'))+1)]
        for key in ('tags','collections'):organization[key]+=d.get(key,[])
    organization={k:sorted(set(v),key=str.casefold) for k,v in organization.items()}
    return jsonify(docs=sorted(docs,key=lambda d:d['created'],reverse=True),templates=templates(),readers=READERS,user=app.extensions['household']['user'](),organization=organization,people=app.extensions['household']['people'](),states=STATES,countries=COUNTRIES)
@app.post('/api/upload')
def upload():
    f=request.files.get('file')
    if not f:abort(400)
    name=secure_filename(f.filename or '')
    if Path(name).suffix.lower() not in READERS:return jsonify(error='This file format is not supported by intake yet. It is not necessarily an executable; keep it for review.'),400
    id=uuid.uuid4().hex;p=ROOT/'inbox'/f'{id}{Path(name).suffix.lower()}';f.save(p)
    failure=app.extensions['intake_reserve'](str(request.form.get('relative_path',f.filename or '')),f.filename or '',p.stat().st_size)
    if failure:
        p.unlink();return jsonify(error=failure[0]),failure[1]
    limit=(512 if p.suffix.lower() in MEDIA_EXTS else 40)*1024*1024
    if p.stat().st_size>limit:
        p.unlink();return jsonify(error='File exceeds its intake limit (40 MB for documents; 512 MB for video/audio).'),413
    hasher=hashlib.sha256()
    with p.open('rb') as stream:
        for block in iter(lambda:stream.read(1024*1024),b''):hasher.update(block)
    digest=hasher.hexdigest()
    with db() as c:
        prior=next((json.loads(r[0]) for r in c.execute('SELECT data FROM docs') if json.loads(r[0]).get('sha256')==digest),None)
    if prior:
        p.unlink();return jsonify(error='Already received: '+(prior['original'] if app.extensions['household']['visible'](prior) else 'an existing document')),409
    source_modified=None
    try:
        milliseconds=int(request.form.get('last_modified',''))
        if 0<=milliseconds<=datetime.now(timezone.utc).timestamp()*1000+86400000:
            source_modified=datetime.fromtimestamp(milliseconds/1000,timezone.utc).isoformat()
    except (ValueError, OverflowError, OSError):pass
    d=dict(size_bytes=p.stat().st_size,source_modified=source_modified,id=id,original=name,path=str(p.relative_to(ROOT)),sha256=digest,created=stamp(),status='review',source_relative_path=str(request.form.get('relative_path',''))[:8192],type='Unclassified',fields={},filename=name,folder='unclassified',text='',notice='',history=[{'at':stamp(),'action':'Received'}])
    d.update(owner_id=app.extensions['household']['user']()['id'],visibility='private',tags=[],collections=[],person_ids=[],batch_id=request.form.get('batch_id'),version=0)
    try:d['text']=extract(p)
    except Exception:d['notice']='Extraction could not finish. Original retained; enter information manually or try another format.'
    save(d);return jsonify(d)
@app.put('/api/docs/<id>')
def edit(id):
    d=getdoc(id);d['_expected_version']=d.get('version',0)
    v=request.get_json() or {};fields=v.get('fields',d.get('fields',{}))
    u=app.extensions['household']['user']()
    if u['role']!='admin' and d.get('owner_id')!=u['id'] and v.get('folder',d['folder'])!=d['folder']:abort(403)
    if not isinstance(fields,dict) or len(fields)>40:abort(400)
    if type(v.get('version',d.get('version',0)))!=int:abort(400)
    if v.get('version',d.get('version',0))!=d.get('version',0):return jsonify(error='Details changed. Reload before saving.'),409
    before={k:d.get(k) for k in ('type','filename','folder','fields','text')}
    for key in ['type','filename','folder','text']:d[key]=str(v.get(key,d[key]))[:250000 if key=='text' else 200]
    d['folder']=archive_folder(d['folder'])
    d['fields']={str(k)[:100]:str(value)[:5000] for k,value in fields.items()}
    template=next((t for t in templates() if t['name']==d['type']),None)
    if d['status']=='archived':
        missing=[f for f in (template or {}).get('required',[]) if not d['fields'].get(f,'').strip()]
        if missing:return jsonify(error='Required: '+', '.join(missing)),400
    for key,schema in (template or {}).get('schema',{}).items():
        value=d['fields'].get(key,'')
        if value and value!=before['fields'].get(key):
            if schema.get('type')=='date':
                try:datetime.strptime(value,'%Y-%m-%d')
                except ValueError:return jsonify(error=key+' must be YYYY-MM-DD.'),400
            if schema.get('type')=='select' and value not in schema.get('options',[]):return jsonify(error='Choose a listed value for '+key),400
    d['version']=d.get('version',0)+1
    changed={k:{'before':before[k],'after':d[k]} for k in before if before[k]!=d[k]}
    d['history'].append({'at':stamp(),'action':'Details edited by '+app.extensions['household']['user']()['name'],'changes':changed})
    save(d);return jsonify(d)
@app.post('/api/docs/<id>/suggest')
def suggest_details(id):
    d=getdoc(id)
    return jsonify(suggest(d.get('text',''),app.extensions['household']['people']()))
@app.post('/api/docs/<id>/approve')
def approve(id):
    d=getdoc(id);d['_expected_version']=d.get('version',0)
    if d['status']=='archived':return jsonify(d)
    template=next((t for t in templates() if t['name']==d['type']),None)
    if not template and d['type']!='Unclassified':return jsonify(error='Choose a known document type.'),400
    missing=[f for f in (template or {}).get('required',[]) if not d['fields'].get(f,'').strip()]
    if missing:return jsonify(error='Required: '+', '.join(missing)),400
    folder=archive_folder(d['folder']);name=secure_filename(d['filename'])
    if not name:return jsonify(error='Enter a filename.'),400
    ext=Path(d['original']).suffix.lower()
    if Path(name).suffix.lower()!=ext:name+=ext
    dest=ROOT/'archive'/folder/id;dest.mkdir(parents=True,exist_ok=True)
    original=confined(d['path']);original.rename(dest/name)
    d.update(path=str((dest/name).relative_to(ROOT)),filename=name,folder=folder,status='archived')
    d['history'].append({'at':stamp(),'action':'Approved and archived'})
    (dest/'record.json').write_text(json.dumps(d,indent=2));(dest/'extracted.txt').write_text(d['text']);save(d);return jsonify(d)
@app.get('/api/docs/<id>/original')
def original(id):
    d=getdoc(id);preview=request.args.get('preview')=='1' and Path(d['original']).suffix.lower() in ['.pdf','.png','.jpg','.jpeg','.webp']
    return send_file(confined(d['path']),download_name=d['original'],as_attachment=not preview)
@app.get('/api/docs/<id>/note')
def note(id):
    d=getdoc(id);text=d['filename']+'\n\n'+'\n'.join(f'{k}: {v}' for k,v in d['fields'].items())+'\n\n'+d['text']
    return send_file(BytesIO(text.encode()),download_name='safenote-draft.txt',as_attachment=True,mimetype='text/plain')
@app.put('/api/templates')
def teach():
    t=request.get_json();name=str(t.get('name','')).strip()[:80];fields=t.get('fields',[]);required=t.get('required',[])
    if not name or not isinstance(fields,list) or not 1<=len(fields)<=40 or not all(isinstance(f,str) and 0<len(f)<=100 for f in fields) or len(set(fields))!=len(fields):abort(400)
    if not isinstance(required,list) or not set(required)<=set(fields):abort(400)
    item=dict(name=name,fields=fields,required=required,folder=archive_folder(str(t.get('folder',name.lower()))))
    existing=next((x for x in templates() if x['name']==name),{})
    schema=t.get('schema',{k:v for k,v in existing.get('schema',{}).items() if k in fields})
    if not isinstance(schema,dict) or not set(schema)<=set(fields):abort(400)
    for f,config in schema.items():
        if not isinstance(config,dict) or config.get('type') not in ('text','date','select','person'):abort(400)
        if config['type']=='select' and (not isinstance(config.get('options'),list) or len(config['options'])>250 or any(not isinstance(x,str) or len(x)>100 for x in config['options'])):abort(400)
    item['schema']=schema
    ts=[x for x in templates() if x['name']!=name]+[item]
    with db() as c:c.execute("UPDATE settings SET data=? WHERE key='templates'",(json.dumps(ts),))
    return jsonify(item)

def archive_folder(value):
    parts=str(value).split('/')
    if len(parts)>8 or any(p in ('','.', '..') or '\\' in p for p in parts):abort(400)
    clean=[secure_filename(p) for p in parts]
    if not all(clean):abort(400)
    return '/'.join(clean)

with db() as c:
    ts=templates()
    for name,fields in [('1098-T Tuition Statement',['Tax year','Student','School','Tuition payments','Scholarships and grants']),('Tax document',['Tax year','Person','Issuer','Form']),('Video',['Title','Recording date','Source','People or event','Notes']),('Audio',['Title','Recording date','Source','Notes'])]:
        if not any(t['name']==name for t in ts):ts.append(dict(name=name,fields=fields,required=['Title'] if name in ('Video','Audio') else fields[:3],folder='videos' if name=='Video' else 'audio' if name=='Audio' else 'taxes'))
    c.execute("UPDATE settings SET data=? WHERE key='templates'",(json.dumps(ts),))

from household import register as register_household
register_household(app,db,stamp,getdoc,save,archive_folder)


from details import STATES,COUNTRIES,BIRTH_FIELDS,BIRTH_SCHEMA,suggest
with db() as c:
    ts=templates();birth=next((t for t in ts if t['name']=='Birth Certificate'),None)
    if not birth:
        birth=dict(name='Birth Certificate',folder='Family_Records/Vital_Records/Birth_Certificates',fields=list(BIRTH_FIELDS),required=['Person']);ts.append(birth)
    else:
        birth['legacy_fields']=[f for f in birth['fields'] if f not in BIRTH_FIELDS]
        birth['fields']=list(dict.fromkeys(BIRTH_FIELDS+birth['fields']))
    birth['schema']=BIRTH_SCHEMA
    c.execute("UPDATE settings SET data=? WHERE key='templates'",(json.dumps(ts),))
from intake import register as register_intake
register_intake(app,db,stamp)

app.config['CATALOG_SNAPSHOT_DIR']=str(STATE/'catalog-snapshots')
from catalog import register
register(app,db,stamp)

if __name__=='__main__':app.run(host='0.0.0.0',port=8080)
